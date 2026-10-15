"""Turns what a user tells us (CV, experience, skills, expectations) into an engine-ready profile."""
import datetime as dt
import io
import re
from functools import lru_cache
from pathlib import Path

import yaml

DATA = Path(__file__).parent / "data"
ROMAN = {1: "i", 2: "ii", 3: "iii", 4: "iv", 5: "v"}
TIERS = ["A", "B", "C", "D"]


@lru_cache(maxsize=1)
def vocab():
    return yaml.safe_load((DATA / "skills.yaml").read_text(encoding="utf-8"))


@lru_cache(maxsize=1)
def regions():
    return yaml.safe_load((DATA / "regions.yaml").read_text(encoding="utf-8"))


def term_re(term):
    return re.compile(r"(?<![a-z0-9])" + re.escape(term.strip().lower()) + r"(?![a-z0-9])")


# ------------------------------------------------------------------ CV reading
def extract_text(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")


MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_RANGE = re.compile(
    r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*[,']?\s*(\d{4})\s*(?:-|–|—|to)\s*"
    r"(?:(present|current|now|till date|ongoing)|(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*[,']?\s*(\d{4}))",
    re.I)


def guess_experience(text: str):
    """Best-effort years of experience: explicit 'N years' statements, else merged Mon-YYYY date ranges."""
    t = text.lower()
    stated = [float(m.group(1)) for m in re.finditer(
        r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s*(?:of\s*)?(?:total\s*|professional\s*|relevant\s*|industry\s*)?(?:experience|exp)", t)]
    stated = [s for s in stated if 0 < s < 40]
    if stated:
        return max(stated)
    now = dt.date.today()
    spans = []
    for m in _RANGE.finditer(text):
        s = int(m.group(2)) * 12 + MONTHS[m.group(1).lower()[:3]]
        e = now.year * 12 + now.month if m.group(3) else int(m.group(5)) * 12 + MONTHS[m.group(4).lower()[:3]]
        if e >= s:
            spans.append((s, e))
    if not spans:
        return None
    spans.sort()
    total, cur_s, cur_e = 0, *spans[0]
    for s, e in spans[1:]:
        if s <= cur_e:
            cur_e = max(cur_e, e)
        else:
            total += cur_e - cur_s
            cur_s, cur_e = s, e
    total += cur_e - cur_s
    return round(total / 12, 1) if total else None


def skill_regexes(name):
    aliases = vocab()["skills"].get(name, [name])
    return [term_re(a) for a in {name, *[str(a) for a in aliases]}]


def analyze_cv(text: str) -> dict:
    t = text.lower()
    counts = {}
    for name in vocab()["skills"]:
        n = sum(len(r.findall(t)) for r in skill_regexes(name))
        if n:
            counts[name] = n
    skills = [k for k, _ in sorted(counts.items(), key=lambda kv: -kv[1])]
    role_hits = {}
    for fam, spec in vocab()["roles"].items():
        n = sum(len(term_re(k).findall(t)) for k in spec["include"])
        if n:
            role_hits[fam] = n
    roles = [k for k, v in sorted(role_hits.items(), key=lambda kv: -kv[1]) if v >= 2][:2]
    return {"skills": skills, "skill_counts": counts, "yoe": guess_experience(text), "roles": roles,
            "chars": len(text)}


# ------------------------------------------------------------------ seniority
def _levels(nums):
    out = []
    for n in nums:
        for stem in ("sde", "sde-", "swe", "engineer", "developer"):
            sep = "" if stem.endswith("-") else " "
            out += [f"{stem}{sep}{n}", f"{stem}{sep}{ROMAN[n]}"]
    return out


def seniority_rules(yoe: float):
    """Returns (exclude_terms, boost_terms) for the title, based on years of experience."""
    ent = ["junior", "associate", "graduate", "entry level", "entry-level", "fresher", "campus", "early career", "new grad"]
    top = ["staff", "principal", "architect", "distinguished", "fellow"]
    if yoe < 1:
        return (["senior", "sr", "lead", "manager", *top, *_levels([2, 3, 4, 5])],
                {**{k: 10 for k in ent}, **{k: 10 for k in _levels([1])}})
    if yoe < 3:
        return (["senior", "sr", "lead", "manager", *top, *_levels([3, 4, 5])],
                {**{k: 6 for k in ent}, **{k: 8 for k in _levels([1, 2])}})
    if yoe < 5:
        return (["lead", "manager", *top, *ent, *_levels([1, 4, 5])],
                {**{k: 8 for k in _levels([2, 3])}, "senior": 4})
    if yoe < 8:
        return (["manager", "distinguished", "fellow", "architect", *ent, *_levels([1, 2])],
                {**{k: 8 for k in _levels([3, 4])}, "senior": 10, "lead": 3, "staff": 3})
    return (["distinguished", "fellow", *ent, *_levels([1, 2, 3])],
            {"senior": 6, "staff": 10, "principal": 8, "lead": 8, "architect": 4, "engineering manager": 3})


# ------------------------------------------------------------------ profile
def _list(v):
    if v is None:
        return []
    if isinstance(v, str):
        return [x.strip() for x in re.split(r"[,\n;]", v) if x.strip()]
    return [str(x).strip() for x in v if str(x).strip()]


def tiers_for_expectation(expected_lpa, country):
    """Indicative: keep tiers whose typical India pay band can reach the expectation."""
    if not expected_lpa or country != "India":
        return list(TIERS)
    bands = regions()["india_tier_bands_lpa"]
    keep = [t for t in TIERS if bands[t][1] >= expected_lpa]
    return keep or ["A"]


def normalize_profile(raw: dict) -> dict:
    v, reg = vocab(), regions()
    yoe = float(raw.get("yoe") or 0)
    country = raw.get("country") or "India"
    if country not in reg["countries"]:
        raise ValueError(f"Unsupported country: {country}")
    cdef = reg["countries"][country]
    canon = {s.lower(): s for s in v["skills"]}

    def canonical(s):
        s = s.strip().lower()
        if s in canon:
            return canon[s]
        for name, aliases in v["skills"].items():
            if s in [str(a).lower() for a in aliases]:
                return name
        return s

    user_skills, seen = [], set()
    for s in _list(raw.get("skills")):
        c = canonical(s)
        if c not in seen:
            seen.add(c)
            user_skills.append(c)
    cv_skills = [canonical(s) for s in _list(raw.get("cv_skills")) if canonical(s) not in seen][:12]
    weights = {s: (10 if i < 3 else 7) for i, s in enumerate(user_skills)}
    weights.update({s: 4 for s in cv_skills})

    roles = [r for r in _list(raw.get("roles")) if r in v["roles"]] or ["backend", "fullstack"]
    include = list(v["generic_software_titles"])
    boost = {}
    for r in roles:
        include += v["roles"][r]["include"]
        boost.update({k: 8 for k in v["roles"][r]["include"] if k not in ("developer", "engineer")})
    include += _list(raw.get("include_keywords"))
    for s in user_skills[:3]:
        boost[s] = 6

    sen_excl, sen_boost = seniority_rules(yoe)
    boost.update(sen_boost)
    exclude = list(v["always_exclude"]) + sen_excl + _list(raw.get("exclude_keywords"))
    for fam, terms in v["family_exclusive_titles"].items():
        if fam == "other" or fam not in roles:
            exclude += terms
    exclude = [e for e in exclude if e.lower() not in {i.lower() for i in include if i not in v["generic_software_titles"]}]

    terms = []
    for r in roles:
        terms += v["roles"][r]["search"]
    terms += [f"{s} developer" for s in user_skills[:2] if s in ("java", "python", "golang", "kotlin", "node.js", "c#", "rust", "scala", "react", "javascript", "typescript", "php", "ruby")]
    terms = list(dict.fromkeys(terms))[:5] or ["software engineer"]

    selected_tiers = [t for t in _list(raw.get("tiers")) if t in TIERS] or list(TIERS)
    exp_lpa = float(raw["expected_lpa"]) if raw.get("expected_lpa") not in (None, "") else None
    allowed = tiers_for_expectation(exp_lpa, country)
    tiers = [t for t in selected_tiers if t in allowed] or selected_tiers

    cities = [c for c in _list(raw.get("cities")) if c in cdef["cities"]]
    return {
        "yoe": yoe, "country": country, "cities": cities, "strict_cities": bool(raw.get("strict_cities", True)),
        "remote_anywhere": bool(raw.get("remote_anywhere", False)), "roles": roles,
        "title_include": include, "title_exclude": exclude, "title_boost": boost,
        "skills": weights, "user_skills": user_skills, "must_have": [canonical(s) for s in _list(raw.get("must_have"))],
        "avoid": [canonical(s) for s in _list(raw.get("avoid"))], "skill_score_cap": 40,
        "tiers": tiers, "expected_lpa": exp_lpa, "search_terms": terms,
        "max_required_years": yoe + float(raw.get("experience_margin", 2)),
        "freshness_days_max": int(raw.get("freshness_days", 45)), "min_score_to_show": int(raw.get("min_score", 35)),
        "apply_now_score": int(raw.get("apply_now_score", 65)), "max_per_company": int(raw.get("max_per_company", 12)),
    }


def load_profile_file(path) -> dict:
    return yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
